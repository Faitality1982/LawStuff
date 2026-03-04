"""
Michigan Index of Authorities Generator
Reads a DOCX brief directly (no LibreOffice needed), extracts all citations
with page numbers, and generates a formatted Index of Authorities in the browser.
"""

import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import threading
import os
import sys
import re
import webbrowser
import tempfile
from collections import defaultdict

try:
    from docx import Document as DocxDocument
    from lxml import etree
except ImportError as e:
    root = tk.Tk(); root.withdraw()
    messagebox.showerror("Missing dependency", str(e))
    sys.exit(1)

REPORTER = r'(?:Mich\.?\s*App\.?|Mich\.?|U\.?S\.?|NW[23]?d|S\.?\s*Ct\.?|SCt\.?)'
IN_RE_PAT = re.compile(r'(?<![A-Za-z])(In\s+re\s+[A-Z][A-Za-z/]+(?:\s+[A-Z][A-Za-z]+)?(?:\s+Minors)?)\s*,\s*(\d+\s+' + REPORTER + r'[^(]{0,150}\(\d{4}\))')
V_IN_RE_PAT = re.compile(r'(?<![A-Za-z])([A-Z][A-Za-z]+(?:\s+[A-Z][A-Za-z]+){0,5}\s+v[s]?\.?\s+[A-Z][A-Za-z]+(?:\s+[A-Z]?[A-Za-z]+){0,5}\s*\(In\s+re\s+[A-Za-z/]+\))\s*,\s*(\d+\s+' + REPORTER + r'[^(]{0,150}\(\d{4}\))')
SIMPLE_V_PAT = re.compile(r'(?<![A-Za-z])([A-Z][A-Za-z]+(?:\s+[A-Z][A-Za-z\.]+){0,3}\s+v[s]?\.?\s+[A-Z][A-Za-z]+(?:\s+[A-Z]?[A-Za-z\.]+){0,3})\s*,\s*(\d+\s+' + REPORTER + r'[^(]{0,150}\(\d{4}\))')
MCL_PAT  = re.compile(r'(MCL\s+\d+[A-Z]?\.\d+[A-Za-z]?(?:\(\d+\))?(?:\([a-z]\))?(?:\([ivxIVX]+\))?)')
MCR_PAT  = re.compile(r'(MCR\s+\d+\.\d+(?:\([A-Za-z0-9]+\))*)')
SKIP_PAT = re.compile(r'TABLE OF CONTENTS|INDEX OF AUTHORITIES|CERTIFICATE OF COMP', re.IGNORECASE)
BODY_PAT = re.compile(r'STATEMENT OF FACTS|STATEMENT OF JURISDICTION|ARGUMENT', re.IGNORECASE)
MCL_SKIP = re.compile(r'MCL\s+7\.\d')
CANON_NAMES = {
    'Family Independence Agency v Boursaw (In re Boursaw)': 'Family Independent Agency v Boursaw (In re Boursaw)',
    'Family Independence Agency v Sours (In re Sours)':    'Family Independent Agency v Sours (In re Sours)',
}

def clean_name(raw):
    raw = raw.strip().rstrip('.,;')
    raw = re.sub(r'^In\s+(?!re\s)', '', raw)
    return CANON_NAMES.get(raw, raw)

def sort_pages(pages):
    roman = {'i':1,'ii':2,'iii':3,'iv':4,'v':5,'vi':6}
    def key(p):
        if p.lower() in roman: return (0, roman[p.lower()])
        try: return (1, int(p))
        except: return (2, p)
    return sorted(pages, key=key)

def _has_pagebreak(para_elem):
    xml = etree.tostring(para_elem, encoding='unicode')
    return 'w:type="page"' in xml or 'lastRenderedPageBreak' in xml

def extract_pages_from_docx(docx_path):
    doc = DocxDocument(docx_path)
    pages_text = []
    current = []
    for para in doc.paragraphs:
        if _has_pagebreak(para._element) and current:
            pages_text.append('\n'.join(current))
            current = []
        text = para.text.strip()
        if text:
            current.append(text)
    if current:
        pages_text.append('\n'.join(current))

    # Fallback: form-feed split
    if len(pages_text) <= 1:
        full = '\n'.join(p.text for p in doc.paragraphs)
        if '\x0c' in full:
            pages_text = full.split('\x0c')

    roman_seq = ['i','ii','iii','iv','v','vi','vii','viii','ix','x']
    roman_idx = arabic_idx = 0
    in_body = False
    labeled = []
    for text in pages_text:
        if not text.strip():
            continue
        if BODY_PAT.search(text):
            in_body = True
        if in_body:
            arabic_idx += 1
            label = str(arabic_idx)
        else:
            label = roman_seq[roman_idx] if roman_idx < len(roman_seq) else str(roman_idx+1)
            roman_idx += 1
        labeled.append((label, text))
    return labeled

def extract_index(docx_path, progress_cb=None):
    if progress_cb: progress_cb("Reading document…")
    pages = extract_pages_from_docx(docx_path)
    if progress_cb: progress_cb(f"Scanning {len(pages)} pages for citations…")

    cases = {}
    statutes = defaultdict(set)
    rules    = defaultdict(set)
    mcl_bare = defaultdict(set)
    in_skip  = False

    for page_label, raw in pages:
        if SKIP_PAT.search(raw):
            in_skip = not bool(BODY_PAT.search(raw))
        if in_skip:
            continue
        flat = re.sub(r'[\n\r]+', ' ', raw)
        flat = re.sub(r'\s+', ' ', flat)
        seen = []
        for pat in [V_IN_RE_PAT, IN_RE_PAT, SIMPLE_V_PAT]:
            for m in pat.finditer(flat):
                if not any(m.start() < e and m.end() > s for s, e in seen):
                    name = clean_name(m.group(1))
                    cite = m.group(2).strip().rstrip('.,;')
                    if name not in cases:
                        cases[name] = {'pages': set(), 'cite': cite}
                    cases[name]['pages'].add(page_label)
                    seen.append((m.start(), m.end()))
        for m in MCL_PAT.finditer(flat):
            val = m.group(1).strip()
            if not MCL_SKIP.match(val):
                if re.match(r'MCL\s+712A\.19b$', val):
                    mcl_bare[val].add(page_label)
                else:
                    statutes[val].add(page_label)
        for m in MCR_PAT.finditer(flat):
            rules[m.group(1).strip()].add(page_label)

    mcl_keys = list(statutes.keys())
    filtered_statutes = {}
    for k in sorted(mcl_keys, key=len, reverse=True):
        is_prefix = any(o != k and o.startswith(k+'(') for o in mcl_keys)
        if not is_prefix:
            filtered_statutes[k] = statutes[k]
        else:
            for o in mcl_keys:
                if o != k and o.startswith(k+'('):
                    statutes[o].update(statutes[k])
    for k, v in mcl_bare.items():
        filtered_statutes[k] = v

    return {
        'cases':    {k: {'pages': sort_pages(v['pages']), 'cite': v['cite']} for k, v in sorted(cases.items())},
        'statutes': {k: sort_pages(v) for k, v in sorted(filtered_statutes.items())},
        'rules':    {k: sort_pages(v) for k, v in sorted(rules.items())},
    }

def build_html(data):
    def ehtml(name, ps, italic=False):
        n = f'<em>{name}</em>' if italic else name
        return f'<div class="entry"><span class="entry-name">{n}</span><span class="dots"></span><span class="pages">{ps}</span></div>\n'
    def ejs(name, ps):
        return f"    ['{name.replace(chr(39), chr(92)+chr(39))}', '{ps}']"

    ch, cj = '', []
    for name, v in data['cases'].items():
        ps = ', '.join(v['pages']); ch += ehtml(name, ps, True); cj.append(ejs(name, ps))
    sh, sj = '', []
    for name, pages in data['statutes'].items():
        ps = ', '.join(pages); sh += ehtml(name, ps); sj.append(ejs(name, ps))
    rh, rj = '', []
    for name, pages in data['rules'].items():
        ps = ', '.join(pages); rh += ehtml(name, ps); rj.append(ejs(name, ps))

    cj_str = ',\n'.join(cj); sj_str = ',\n'.join(sj); rj_str = ',\n'.join(rj)

    return f"""<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8">
<title>Index of Authorities</title>
<style>
*{{box-sizing:border-box;margin:0;padding:0}}
body{{font-family:Arial,sans-serif;font-size:12pt;max-width:740px;margin:30px auto;padding:20px 30px 50px;color:#000;background:#fff}}
h1{{text-align:center;font-size:13pt;font-weight:bold;text-decoration:underline;margin-bottom:24px;letter-spacing:.04em}}
h2{{font-size:12pt;font-weight:normal;text-decoration:underline;margin:20px 0 8px}}
.entry{{display:flex;align-items:baseline;margin-bottom:5px;line-height:1.5}}
.entry-name{{flex-shrink:0;max-width:75%}}
.entry-name em{{font-style:italic}}
.dots{{flex:1;border-bottom:1px dotted #000;margin:0 4px 3px;min-width:20px}}
.pages{{flex-shrink:0;white-space:nowrap}}
.copy-btn{{display:block;margin:0 0 24px auto;padding:7px 18px;background:#1a3a6b;color:#fff;border:none;border-radius:4px;font-size:11pt;cursor:pointer;font-family:Arial,sans-serif}}
.copy-btn:hover{{background:#0f2548}}
.copy-btn.copied{{background:#2a7a3b}}
@media print{{.copy-btn{{display:none}}}}
</style></head><body>
<button class="copy-btn" id="copyBtn" onclick="copyIndex()">&#128203; Copy as Plain Text</button>
<h1>INDEX OF AUTHORITIES</h1>
<h2>CASE LAW</h2>
{ch}
<h2>STATUTES &amp; OTHER AUTHORITIES</h2>
{sh}
<h2>MICHIGAN COURT RULES</h2>
{rh}
<script>
const CASES=[{cj_str}];
const STATS=[{sj_str}];
const RULES=[{rj_str}];
function fmt(n,p){{return n+'.'.repeat(Math.max(3,70-n.length-p.length))+p;}}
function copyIndex(){{
  const L=['INDEX OF AUTHORITIES','','CASE LAW',''];
  CASES.forEach(([n,p])=>L.push(fmt(n,p)));
  L.push('','STATUTES & OTHER AUTHORITIES','');
  STATS.forEach(([n,p])=>L.push(fmt(n,p)));
  L.push('','MICHIGAN COURT RULES','');
  RULES.forEach(([n,p])=>L.push(fmt(n,p)));
  navigator.clipboard.writeText(L.join('\\n')).then(()=>{{
    const b=document.getElementById('copyBtn');
    b.textContent='✓ Copied!';b.classList.add('copied');
    setTimeout(()=>{{b.textContent='📋 Copy as Plain Text';b.classList.remove('copied');}},2500);
  }}).catch(()=>{{const t=document.createElement('textarea');t.value=L.join('\\n');
    document.body.appendChild(t);t.select();document.execCommand('copy');document.body.removeChild(t);alert('Copied!');}});
}}
</script></body></html>"""

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Michigan Index of Authorities Generator")
        self.geometry("560x300")
        self.resizable(False, False)
        self.configure(bg="#f4f6fb")
        self._build_ui()

    def _build_ui(self):
        tk.Label(self, text="Michigan Index of Authorities",
                 font=("Arial",15,"bold"), bg="#1a3a6b", fg="white", pady=14).pack(fill="x")
        frame = tk.Frame(self, bg="#f4f6fb", padx=30, pady=20)
        frame.pack(fill="both", expand=True)
        tk.Label(frame, text="Select your appellate brief (.docx):",
                 font=("Arial",11), bg="#f4f6fb").pack(anchor="w")
        row = tk.Frame(frame, bg="#f4f6fb"); row.pack(fill="x", pady=8)
        self.path_var = tk.StringVar()
        tk.Entry(row, textvariable=self.path_var, font=("Arial",10), width=42).pack(side="left", padx=(0,8))
        tk.Button(row, text="Browse…", command=self._browse,
                  bg="#1a3a6b", fg="white", relief="flat", font=("Arial",10), padx=10).pack(side="left")
        self.status = tk.StringVar(value="Ready.")
        tk.Label(frame, textvariable=self.status, font=("Arial",9), fg="#555", bg="#f4f6fb").pack(anchor="w")
        self.progress = ttk.Progressbar(frame, mode="indeterminate", length=480)
        self.progress.pack(pady=(6,12))
        tk.Button(frame, text="Generate Index of Authorities", command=self._run,
                  bg="#1a3a6b", fg="white", font=("Arial",12,"bold"), relief="flat", padx=20, pady=8).pack()

    def _browse(self):
        path = filedialog.askopenfilename(title="Select DOCX Brief",
            filetypes=[("Word Documents","*.docx"),("All files","*.*")])
        if path: self.path_var.set(path)

    def _run(self):
        path = self.path_var.get().strip()
        if not path or not os.path.exists(path):
            messagebox.showerror("No file", "Please select a valid .docx file first.")
            return
        self.progress.start(12); self.status.set("Working…")
        threading.Thread(target=self._worker, args=(path,), daemon=True).start()

    def _worker(self, path):
        try:
            data = extract_index(path, progress_cb=lambda s: self.status.set(s))
            html = build_html(data)
            out  = tempfile.NamedTemporaryFile(suffix=".html", delete=False,
                       mode="w", encoding="utf-8", prefix="index_of_authorities_")
            out.write(html); out.close()
            self.after(0, lambda p=out.name, d=data: self._done(p, d))
        except Exception as ex:
            msg = str(ex)
            self.after(0, lambda m=msg: self._error(m))

    def _done(self, html_path, data):
        self.progress.stop()
        nc, ns, nr = len(data['cases']), len(data['statutes']), len(data['rules'])
        self.status.set(f"Done!  {nc} cases · {ns} statutes · {nr} rules")
        webbrowser.open(f"file:///{html_path}")
        messagebox.showinfo("Index Generated",
            f"Index of Authorities opened in your browser.\n\n"
            f"  Cases:    {nc}\n  Statutes: {ns}\n  Rules:    {nr}\n\n"
            f"Use the 'Copy as Plain Text' button to paste into Word.")

    def _error(self, msg):
        self.progress.stop(); self.status.set("Error — see dialog.")
        messagebox.showerror("Error", msg)

if __name__ == "__main__":
    App().mainloop()
